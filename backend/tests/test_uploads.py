import io

from docx import Document as DocxDocument
from fpdf import FPDF

from app.core.config import get_settings
from tests.helpers import API, REACTOR_TEXT, db_scalar, register_and_login


def make_pdf(text: str | None) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    if text:
        pdf.set_font("Helvetica", size=12)
        pdf.multi_cell(0, 8, text)
    return bytes(pdf.output())


def make_docx() -> bytes:
    document = DocxDocument()
    document.add_paragraph("Quillham station overview.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Coolant"
    table.cell(0, 1).text = "liquid helium"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def upload(client, headers, filename, content, title=None):
    return await client.post(
        f"{API}/documents/upload",
        files={"file": (filename, content, "application/octet-stream")},
        data={"title": title} if title is not None else None,
        headers=headers,
    )


async def test_txt_upload_creates_document_titled_after_the_file(client):
    headers = await register_and_login(client, "a@example.com")
    response = await upload(client, headers, "../../Reactor notes.txt", REACTOR_TEXT.encode())
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["title"] == "Reactor notes"
    assert body["chunk_count"] >= 1


async def test_markdown_upload_with_explicit_title(client):
    headers = await register_and_login(client, "a@example.com")
    response = await upload(
        client, headers, "notes.md", b"# Heading\n\nSome text here.", title="My Title"
    )
    assert response.status_code == 201, response.text
    assert response.json()["title"] == "My Title"


async def test_pdf_text_is_extracted_and_indexed(client):
    headers = await register_and_login(client, "a@example.com")
    response = await upload(client, headers, "reactor.pdf", make_pdf(REACTOR_TEXT))
    assert response.status_code == 201, response.text
    stored = await db_scalar("SELECT string_agg(content, ' ') FROM chunks")
    assert "Quillham" in stored


async def test_docx_paragraphs_and_tables_are_extracted(client):
    headers = await register_and_login(client, "a@example.com")
    response = await upload(client, headers, "station.docx", make_docx())
    assert response.status_code == 201, response.text
    stored = await db_scalar("SELECT string_agg(content, ' ') FROM chunks")
    assert "Quillham station overview." in stored
    assert "Coolant | liquid helium" in stored


async def test_pdf_without_text_is_rejected(client):
    headers = await register_and_login(client, "a@example.com")
    response = await upload(client, headers, "scan.pdf", make_pdf(None))
    assert response.status_code == 422
    assert "No readable text" in response.json()["detail"]


async def test_unsupported_file_type_is_rejected(client):
    headers = await register_and_login(client, "a@example.com")
    response = await upload(client, headers, "run.exe", b"MZ\x90\x00")
    assert response.status_code == 415


async def test_non_utf8_text_file_is_rejected(client):
    headers = await register_and_login(client, "a@example.com")
    response = await upload(client, headers, "notes.txt", b"\xff\xfe\xfa")
    assert response.status_code == 422
    assert "UTF-8" in response.json()["detail"]


async def test_oversized_file_is_rejected(client, monkeypatch):
    monkeypatch.setattr(get_settings(), "MAX_UPLOAD_BYTES", 100)
    headers = await register_and_login(client, "a@example.com")
    response = await upload(client, headers, "big.txt", b"a" * 101)
    assert response.status_code == 413


async def test_upload_requires_login(client):
    response = await upload(client, {}, "notes.txt", b"hello")
    assert response.status_code == 401


async def test_fake_pdf_is_rejected(client):
    headers = await register_and_login(client, "a@example.com")
    response = await upload(client, headers, "fake.pdf", b"hello world")
    assert response.status_code == 422
    assert "not a valid PDF" in response.json()["detail"]